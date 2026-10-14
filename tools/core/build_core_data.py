"""Build the /core page data (core/core_data.json) and the Excel file (core/CORE_list.xlsx) from SEC EDGAR only.

Inputs: core/screen_sec.csv (tools/core/screen_sec.py), core/sec/<T>.json (tools/core/sec_facts.py),
core/review/<T>.json (CORE review), core/current_state.json (risk tier, Cheap Score), core/CORE_RULES.md.
Run: python tools/core/build_core_data.py [repo_root]
"""
import json, os, sys, glob, math, datetime
TODAY_S = datetime.date.today().isoformat()
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

ROOT = sys.argv[1] if len(sys.argv) > 1 else '.'
P = lambda *a: os.path.join(ROOT, 'core', *a)
S = pd.read_csv(P('screen_sec.csv'))
state = json.load(open(P('current_state.json')))
ST = {r['symbol']: r for r in state['rows']}
REV = {}
for f in glob.glob(P('review', '*.json')):
    j = json.load(open(f)); REV[j['ticker']] = j
SEC = {os.path.basename(f)[:-5]: json.load(open(f)) for f in glob.glob(P('sec', '*.json')) if not os.path.basename(f).startswith('_')}
num = lambda v: None if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))) else (float(v) if not isinstance(v, (int, str, bool)) else v)
PF = lambda b: 'PASS' if b is True or b == 'True' else 'FAIL' if b is False or b == 'False' else None
HEB = {'profit': 'רווחיות', 'cash': 'תזרים', 'returns': 'תשואה על ההון', 'balance': 'מאזן', 'growth': 'צמיחה', 'dilution': 'דילול'}
CAP = {'LOW': .10, 'MEDIUM': .10, 'HIGH': .06, 'EXTREME': .03}
order = {'CORE-ELIGIBLE': 0, 'WATCH': 1, 'NON-CORE': 2}
S = S.sort_values(['label', 'ticker'], key=lambda s: s.map(order) if s.name == 'label' else s)

# ---------------- List
lcols = ['סימול', 'שם', 'מחלקה', 'תווית', 'מספר כישלונות', '1 רווחיות', '2 תזרים', '3 תשואה על ההון', '4 מאזן', '5 צמיחה', '6 דילול',
         'סיבת הכישלון', 'תקופה אחרונה', 'Risk Tier', 'תקרת פוזיציה CORE', 'Cheap Score', 'רבעון זול', 'SMA200 עולה', 'Cheap Score שמיש לשיטה',
         'Quality סופי (סקירה)', 'הערות נתונים']
lfmt = ['txt'] * 14 + ['pct', 'int'] + ['txt'] * 5
lrows = []
for _, r in S.iterrows():
    t = r.ticker; st = ST.get(t, {}); fails = [x for x in str(r.fails).split(';') if x and x != 'nan']
    q = st.get('Cheap_Quartile'); rise = st.get('sma200_rising')
    lrows.append([t, SEC.get(t, {}).get('name'), r.cls, r.label, len(fails)] + [PF(r[f'pass_{k}']) if r.cls == 'STD' or k not in ('cash', 'balance') or r.cls == 'UTIL/REIT' else ('n/a' if r.cls == 'FIN' else PF(r[f'pass_{k}'])) for k in HEB]
                 + ['; '.join(HEB[x] for x in fails), f"{r.latest_kind} עד {r.latest_end}", st.get('Risk_Tier'), CAP.get(st.get('Risk_Tier')),
                    num(st.get('Cheap_Score')), q, 'כן' if rise else 'לא', q if (q and rise) else 'לא רלוונטי (SMA200 יורד)' if q else None,
                    (REV.get(t) or {}).get('final_quality', 'טרם נבדקה'), r.note if isinstance(r.note, str) else ''])

# ---------------- Data (metrics)
dspec = [('ticker', 'סימול', 'txt'), ('cls', 'מחלקה', 'txt'), ('latest_kind', 'סוג תקופה', 'txt'), ('latest_end', 'סוף התקופה', 'txt'), ('data_status', 'מצב הנתונים', 'txt'), ('balance_date', 'תאריך מאזן', 'txt'), ('n_fy', 'שנים בנתונים', 'int'),
         ('revenue', 'הכנסות (מיליון)', 'int'), ('rev_cagr3', 'צמיחה שנתית 3 שנים', 'pct'), ('rev_yoy', 'צמיחה מול שנה קודמת', 'pct'), ('op_margin', 'שולי רווח תפעולי', 'pct'),
         ('net_income', 'רווח נקי (מיליון)', 'int'), ('ni_pos_years', 'שנים עם רווח', 'int'), ('fcf', 'FCF (מיליון)', 'int'), ('fcf_pos_years', 'שנים עם FCF חיובי', 'int'),
         ('ocf_pos_years', 'שנים עם תזרים מפעילות חיובי', 'int'), ('roic', 'ROIC', 'pct'), ('roe', 'ROE', 'pct'), ('net_debt', 'חוב נטו (מיליון)', 'int'),
         ('nd_ebitda', 'חוב נטו / EBITDA', 'x'), ('int_cov', 'כיסוי ריבית', 'x'), ('share_chg3', 'שינוי מניות 3 שנים', 'pct'), ('sbc_rev', 'תגמול במניות / הכנסות', 'pct'), ('note', 'הערות', 'txt')]
MIL = {'revenue', 'net_income', 'fcf', 'net_debt'}
drows = [[(num(r[k]) / 1e6 if k in MIL and num(r[k]) is not None else (r[k] if isinstance(r[k], str) else num(r[k]))) for k, _, _ in dspec] for _, r in S.iterrows()]

# ---------------- Raw (EDGAR periods)
rk = ['revenue', 'operating_income', 'pretax_income', 'net_income', 'eps_diluted', 'shares_diluted', 'operating_cash_flow', 'capex', 'sbc', 'interest_expense', 'dna']
rcols = ['ticker', 'סוג', 'התחלה', 'סוף'] + rk + ['תגית הכנסות (EDGAR)']
rrows = []
for t in S.ticker:
    c = SEC.get(t, {}).get('concepts', {})
    periods = {}
    for k in rk:
        for kind in ('annual', 'quarters'):
            for x in (c.get(k) or {}).get(kind, []):
                if x['end'] < '2021-01-01': continue
                periods.setdefault((kind, x['start'], x['end']), {})[k] = x['val']
    for (kind, s0, e0), vals in sorted(periods.items(), key=lambda kv: (kv[0][2], kv[0][0]), reverse=True):
        rrows.append([t, 'שנה' if kind == 'annual' else 'רבעון', s0, e0] + [vals.get(k) for k in rk] + [(c.get('revenue') or {}).get('tag')])

# ---------------- Review
rvspec = [('ticker', 'סימול'), ('excel_status', 'סטטוס מסנן בזמן הסקירה'), ('final_quality', 'Quality סופי'), ('override', 'שינוי?'), ('reason', 'הסיבה'),
          ('valuation', 'שווי'), ('cyclicality', 'שיא מחזורי?'), ('red_flags', 'דגלים אדומים'), ('business_quality', 'איכות העסק'), ('growth', 'צמיחה'),
          ('profitability', 'רווחיות'), ('fcf', 'תזרים'), ('balance_sheet', 'מאזן'), ('capital_allocation', 'הקצאת הון'), ('structural_risks', 'סיכונים מבניים'), ('sources', 'מקורות')]
rv = []
for t, j in sorted(REV.items(), key=lambda kv: ({'CORE': 0, 'WATCH': 1, 'NON-CORE': 2}.get(kv[1]['final_quality'], 3), kv[0])):
    row = []
    for k, _ in rvspec:
        v = j.get(k)
        if k == 'override': v = 'כן' if v else 'לא'
        if k == 'excel_status' and v == 'CORE': v = 'CORE-ELIGIBLE'
        if k == 'sources': v = ' | '.join(f"{s.get('title', '')} ({s.get('date', '')}) {s.get('url', '')}" for s in (v or []))
        row.append(v)
    rv.append(row)

rules_text = open(P('CORE_RULES.md')).read()
# ---------------- Quarters: every quarter EDGAR has (last 12), with year-over-year growth and trailing-4-quarter sums
qcols = ['סימול', 'סוף רבעון', 'הכנסות (M)', 'צמיחה מול רבעון מקביל', 'רווח תפעולי (M)', 'מרווח תפעולי', 'רווח נקי (M)', 'EPS מדולל', 'FCF (M)', 'הכנסות 4 רבעונים (M)', 'מרווח תפעולי 4 רבעונים', 'רווח נקי 4 רבעונים (M)', 'FCF 4 רבעונים (M)', 'צמיחה 4 רבעונים מול שנה קודמת']
qfmt = ['txt', 'txt', 'int', 'pct', 'int', 'pct', 'int', 'dec', 'int', 'int', 'pct', 'int', 'int', 'pct']
qrows = []
for t in sorted(SEC):
    c = SEC[t].get('concepts') or {}
    Q = {k: {x['end']: x['val'] for x in (c.get(k) or {}).get('quarters', []) if x['end'] <= TODAY_S} for k in ('revenue', 'operating_income', 'net_income', 'eps_diluted', 'operating_cash_flow', 'capex')}
    ends = sorted(Q['revenue'])
    def near(dct, e, days):
        import datetime as dd
        best = None
        for k2 in dct:
            g = abs((dd.date.fromisoformat(e) - dd.date.fromisoformat(k2)).days - days)
            if g <= 20 and (best is None or g < best[0]): best = (g, k2)
        return best[1] if best else None
    def ttm(k, e):
        es = sorted(x for x in Q[k] if x <= e)[-4:]
        if len(es) < 4 or es[-1] != e: return None
        import datetime as dd
        if (dd.date.fromisoformat(es[-1]) - dd.date.fromisoformat(es[0])).days > 300: return None
        return sum(Q[k][x] for x in es)
    for e in ends[-12:][::-1]:
        rv_, oi_, ni_, cf_, cp_ = Q['revenue'].get(e), Q['operating_income'].get(e), Q['net_income'].get(e), Q['operating_cash_flow'].get(e), Q['capex'].get(e)
        py = near(Q['revenue'], e, 364)
        tr, to, tn, tc, tp = ttm('revenue', e), ttm('operating_income', e), ttm('net_income', e), ttm('operating_cash_flow', e), ttm('capex', e)
        tr1 = ttm('revenue', py) if py else None
        M = lambda v: v / 1e6 if v is not None else None
        qrows.append([t, e, M(rv_), (rv_ / Q['revenue'][py] - 1) if py and rv_ and Q['revenue'][py] > 0 else None, M(oi_), (oi_ / rv_) if oi_ is not None and rv_ else None,
                      M(ni_), Q['eps_diluted'].get(e), M(cf_ - abs(cp_)) if cf_ is not None and cp_ is not None else None, M(tr), (to / tr) if to is not None and tr else None, M(tn),
                      M(tc - abs(tp)) if tc is not None and tp is not None else None, (tr / tr1 - 1) if tr and tr1 and tr1 > 0 else None])
import datetime as _dt
fetched = sorted({j.get('fetched') for j in SEC.values() if j.get('fetched')})
# Quality Master: one row per company (EDGAR screen + business review), the file ChatGPT asked for
qm_cols = ['symbol', 'name', 'EDGAR_label', 'EDGAR_fails', 'EDGAR_period', 'data_status', 'EDGAR_fetched', 'Review_final', 'review_flag', 'valuation_status', 'review_date',
           'cyclicality', 'thesis', 'override', 'override_reason', 'red_flags', 'structural_risks', 'growth_note', 'profitability_note', 'fcf_note', 'balance_note', 'capital_allocation_note',
           'revenue_m', 'rev_cagr3', 'rev_yoy', 'op_margin', 'net_income_m', 'fcf_m', 'roic', 'roe', 'nd_ebitda', 'int_cov', 'share_chg3', 'sbc_rev', 'data_notes', 'sources']
qm = []
for _, r in S.iterrows():
    t = r.ticker; j = REV.get(t) or {}
    seen = (j.get('excel_status') or '').replace('CORE-ELIGIBLE', 'CORE'); lab = str(r.label).replace('CORE-ELIGIBLE', 'CORE')
    flag = 'REVIEW REQUIRED' if seen and seen != lab else ''
    m = lambda k: num(r[k])
    qm.append([t, SEC.get(t, {}).get('name'), r.label, '; '.join(HEB.get(x, x) for x in str(r.fails).split(';') if x and x != 'nan'), f"{r.latest_kind} {r.latest_end}", r.get('data_status') if isinstance(r.get('data_status'), str) else None,
               SEC.get(t, {}).get('fetched'), j.get('final_quality', 'NOT REVIEWED'), flag, (j.get('valuation') or '').split(' ')[0].split('—')[0].strip() or None,
               j.get('review_date', '2026-10-13'), j.get('cyclicality'), j.get('business_quality'), 'yes' if j.get('override') else 'no', j.get('reason') if j.get('override') else '',
               j.get('red_flags'), j.get('structural_risks'), j.get('growth'), j.get('profitability'), j.get('fcf'), j.get('balance_sheet'), j.get('capital_allocation'),
               (m('revenue') or 0) / 1e6 if m('revenue') is not None else None, m('rev_cagr3'), m('rev_yoy'), m('op_margin'), (m('net_income') / 1e6) if m('net_income') is not None else None,
               (m('fcf') / 1e6) if m('fcf') is not None else None, m('roic'), m('roe'), m('nd_ebitda'), m('int_cov'), m('share_chg3'), m('sbc_rev'),
               r.note if isinstance(r.note, str) else '', ' | '.join(f"{x.get('title', '')} ({x.get('date', '')}) {x.get('url', '')}" for x in (j.get('sources') or []))])
out = {'source': 'SEC EDGAR', 'screen_date': state.get('date'), 'built': _dt.datetime.utcnow().strftime('%Y-%m-%d %H:%MZ'), 'edgar_fetched': fetched[-1] if fetched else None, 'reviews_date': max((j.get('review_date') or '2026-10-13') for j in REV.values()),
       'quality_master': {'cols': qm_cols, 'rows': qm},
       'list': {'cols': lcols, 'fmt': lfmt, 'rows': lrows},
       'review': {'cols': [b for _, b in rvspec], 'fmt': ['txt'] * len(rvspec), 'rows': rv},
       'data': {'cols': [b for _, b, _ in dspec], 'fmt': [f for _, _, f in dspec], 'rows': drows},
       'quarters': {'cols': qcols, 'fmt': qfmt, 'rows': qrows},
       'raw': {'cols': rcols, 'fmt': ['txt'] * 4 + ['int'] * 4 + ['dec'] + ['int'] * 6 + ['txt'], 'rows': rrows},
       'rules': [[line.lstrip('#- ').strip(), None] for line in rules_text.splitlines() if line.strip()]}
json.dump(out, open(P('core_data.json'), 'w'), ensure_ascii=False, allow_nan=False, default=str)

# ---------------- Excel
wb = Workbook(); AR = 'Arial'
hf = Font(name=AR, bold=True, color='FFFFFF'); hfill = PatternFill('solid', fgColor='1F3864'); base = Font(name=AR, size=10)
NF = {'pct': '0.0%', 'int': '#,##0;(#,##0);-', 'x': '0.0x', 'dec': '0.00', 'txt': None}
def sheet(ws, cols, fmts, rows, wide=()):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=1, column=i, value=c); x.font = hf; x.fill = hfill; x.alignment = Alignment(wrap_text=True, vertical='center')
    for ri, r in enumerate(rows, 2):
        for ci, v in enumerate(r, 1):
            x = ws.cell(row=ri, column=ci, value=v); x.font = base
            if NF.get(fmts[ci - 1]) and isinstance(v, (int, float)): x.number_format = NF[fmts[ci - 1]]
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(rows) + 1}'
    for i, c in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = 40 if c in wide else 13
L = wb.active; L.title = 'List'; sheet(L, lcols, lfmt, lrows, wide=('שם', 'סיבת הכישלון', 'הערות נתונים', 'Cheap Score שמיש לשיטה'))
fill = lambda c: PatternFill('solid', fgColor=c)
n = len(lrows) + 1
for v, c in (('"CORE-ELIGIBLE"', 'C6EFCE'), ('"WATCH"', 'FFEB9C'), ('"NON-CORE"', 'F2F2F2')): L.conditional_formatting.add(f'D2:D{n}', CellIsRule(operator='equal', formula=[v], fill=fill(c)))
L.conditional_formatting.add(f'F2:K{n}', CellIsRule(operator='equal', formula=['"FAIL"'], fill=fill('FFC7CE')))
RVs = wb.create_sheet('Review'); sheet(RVs, [b for _, b in rvspec], ['txt'] * len(rvspec), rv, wide=tuple(b for _, b in rvspec[4:]))
Ds = wb.create_sheet('Data'); sheet(Ds, [b for _, b, _ in dspec], [f for _, _, f in dspec], drows, wide=('הערות',))
Qs = wb.create_sheet('Quarters'); sheet(Qs, qcols, qfmt, qrows)
Rs = wb.create_sheet('Rules'); Rs.sheet_view.rightToLeft = True
for i, line in enumerate([l for l in rules_text.splitlines() if l.strip()], 1):
    x = Rs.cell(row=i, column=1, value=line.lstrip('#- ').strip()); x.font = Font(name=AR, size=10, bold=line.startswith('#')); x.alignment = Alignment(wrap_text=True)
Rs.column_dimensions['A'].width = 120
Ws = wb.create_sheet('Raw_EDGAR'); sheet(Ws, rcols, out['raw']['fmt'], rrows, wide=('תגית הכנסות (EDGAR)',))
for sh in (L, Ds, Ws):
    x = sh.cell(row=sh.max_row + 2, column=1, value='מקור: SEC EDGAR (companyfacts). תוויות מחושבות ב-tools/core/screen_sec.py לפי core/CORE_RULES.md. Risk Tier ו-Cheap Score מהנרות שלנו, נכון ל-' + str(state.get('date')))
    x.font = Font(name=AR, italic=True, size=9)
wb.save(P('CORE_list.xlsx'))
print('core_data.json + CORE_list.xlsx:', len(lrows), 'stocks,', len(rrows), 'raw periods,', len(rv), 'reviews; labels', S.label.value_counts().to_dict())
